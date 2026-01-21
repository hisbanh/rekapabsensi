"""
Unit tests for teacher report service.

Tests the teacher report service functionality including PDF and Excel generation.
"""

from django.test import TestCase
from django.contrib.auth.models import User
from django.utils import timezone
from datetime import timedelta
from io import BytesIO
from uuid import UUID

from openpyxl import load_workbook

from attendance.models import (
    Teacher, TeacherAttendance, TeacherSchedule, 
    Subject, Classroom, AcademicLevel
)
from attendance.services.teacher_report_service import TeacherReportService, TeacherReportServiceError


class TeacherReportServiceExcelTestCase(TestCase):
    """Test cases for Excel export functionality."""
    
    def setUp(self):
        """Set up test data."""
        # Create admin user
        self.admin_user = User.objects.create_superuser(
            username='admin',
            email='admin@test.com',
            password='testpass123'
        )
        
        # Create academic level
        self.academic_level, _ = AcademicLevel.objects.get_or_create(
            code='SMP',
            defaults={
                'name': 'Sekolah Menengah Pertama',
                'level_type': 'SMP',
                'min_grade': 7,
                'max_grade': 9
            }
        )
        
        # Create classroom
        self.classroom = Classroom.objects.create(
            name='8A',
            academic_level=self.academic_level,
            grade=8,
            section='A',
            is_active=True
        )
        
        # Create subjects
        self.subject1 = Subject.objects.create(
            code='MAT',
            name='Matematika',
            category='UMUM',
            is_active=True
        )
        
        self.subject2 = Subject.objects.create(
            code='FIS',
            name='Fisika',
            category='UMUM',
            is_active=True
        )
        
        # Create teachers
        self.teacher1 = Teacher.objects.create(
            nip='1234567890',
            full_name='Ahmad Yusuf',
            email='ahmad@test.com',
            employment_date=timezone.now().date() - timedelta(days=365),
            employment_status='ACTIVE',
            is_active=True
        )
        
        self.teacher2 = Teacher.objects.create(
            nip='0987654321',
            full_name='Fatimah Zahra',
            email='fatimah@test.com',
            employment_date=timezone.now().date() - timedelta(days=180),
            employment_status='ACTIVE',
            is_active=True
        )
        
        # Create schedules
        self.schedule1 = TeacherSchedule.objects.create(
            teacher=self.teacher1,
            subject=self.subject1,
            classroom=self.classroom,
            day_of_week=0,  # Monday
            jp_start=1,
            jp_end=2,
            effective_date=timezone.now().date() - timedelta(days=30),
            is_active=True
        )
        
        self.schedule2 = TeacherSchedule.objects.create(
            teacher=self.teacher2,
            subject=self.subject2,
            classroom=self.classroom,
            day_of_week=1,  # Tuesday
            jp_start=3,
            jp_end=4,
            effective_date=timezone.now().date() - timedelta(days=30),
            is_active=True
        )
        
        # Create attendance records
        today = timezone.now().date()
        statuses = ['HADIR', 'HADIR', 'HADIR', 'SAKIT', 'IZIN']
        
        for i in range(5):
            date = today - timedelta(days=i)
            
            # Teacher 1 attendance
            TeacherAttendance.objects.create(
                teacher=self.teacher1,
                schedule=self.schedule1,
                date=date,
                jp_number=1,
                status=statuses[i],
                recorded_by=self.admin_user,
                notes=f'Test note {i}'
            )
            
            # Teacher 2 attendance
            TeacherAttendance.objects.create(
                teacher=self.teacher2,
                schedule=self.schedule2,
                date=date,
                jp_number=3,
                status='HADIR',
                recorded_by=self.admin_user
            )
    
    def test_export_excel_basic(self):
        """Test basic Excel export functionality."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date
        )
        
        # Verify it returns bytes
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 0)
    
    def test_export_excel_has_required_sheets(self):
        """Test that Excel export contains required sheets."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date
        )
        
        # Load workbook
        wb = load_workbook(BytesIO(excel_bytes))
        
        # Check required sheets exist
        self.assertIn('Ringkasan', wb.sheetnames)
        self.assertIn('Detail Kehadiran', wb.sheetnames)
    
    def test_export_excel_summary_sheet_structure(self):
        """Test that summary sheet has correct structure."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date
        )
        
        # Load workbook
        wb = load_workbook(BytesIO(excel_bytes))
        ws = wb['Ringkasan']
        
        # Check header
        self.assertEqual(ws['A1'].value, 'RINGKASAN ABSENSI USTADZ')
        
        # Check overall statistics section
        self.assertEqual(ws['A4'].value, 'Statistik Keseluruhan')
        self.assertEqual(ws['A5'].value, 'Keterangan')
        self.assertEqual(ws['B5'].value, 'Jumlah')
        
        # Check teacher performance section
        self.assertEqual(ws['A16'].value, 'Performa Per Ustadz')
        
        # Check teacher performance headers
        expected_headers = ['No', 'NIP', 'Nama', 'Total JP', 'Hadir', 'Sakit', 'Izin', 'Cuti', 'Dinas', 'Alpa', 'Persentase']
        for col_idx, expected_header in enumerate(expected_headers, 1):
            actual_value = ws.cell(row=17, column=col_idx).value
            self.assertEqual(actual_value, expected_header)
    
    def test_export_excel_detail_sheet_structure(self):
        """Test that detail sheet has correct structure."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date
        )
        
        # Load workbook
        wb = load_workbook(BytesIO(excel_bytes))
        ws = wb['Detail Kehadiran']
        
        # Check headers
        expected_headers = ['No', 'Tanggal', 'NIP', 'Nama Ustadz', 'JP', 'Status', 'Mata Pelajaran', 'Kelas', 'Keterangan', 'Dicatat Oleh', 'Waktu Catat']
        for col_idx, expected_header in enumerate(expected_headers, 1):
            actual_value = ws.cell(row=1, column=col_idx).value
            self.assertEqual(actual_value, expected_header)
    
    def test_export_excel_with_teacher_ids(self):
        """Test Excel export with specific teacher IDs creates per-teacher sheets."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date,
            teacher_ids=[self.teacher1.id, self.teacher2.id]
        )
        
        # Load workbook
        wb = load_workbook(BytesIO(excel_bytes))
        
        # Check that per-teacher sheets exist
        self.assertIn('Ahmad Yusuf', wb.sheetnames)
        self.assertIn('Fatimah Zahra', wb.sheetnames)
    
    def test_export_excel_per_teacher_sheet_structure(self):
        """Test that per-teacher sheets have correct structure."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date,
            teacher_ids=[self.teacher1.id]
        )
        
        # Load workbook
        wb = load_workbook(BytesIO(excel_bytes))
        ws = wb['Ahmad Yusuf']
        
        # Check header
        self.assertIn('LAPORAN ABSENSI', ws['A1'].value)
        self.assertIn('Ahmad Yusuf', ws['A1'].value)
        
        # Check NIP in header
        self.assertIn('1234567890', ws['A2'].value)
        
        # Check summary section
        self.assertEqual(ws['A4'].value, 'Ringkasan')
        self.assertEqual(ws['A5'].value, 'Status')
        
        # Check detail section
        self.assertEqual(ws['A14'].value, 'Detail Kehadiran')
    
    def test_export_excel_has_formulas(self):
        """Test that Excel export contains formulas for calculations."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date,
            teacher_ids=[self.teacher1.id]
        )
        
        # Load workbook
        wb = load_workbook(BytesIO(excel_bytes))
        ws = wb['Ahmad Yusuf']
        
        # Check that percentage cells contain formulas
        # Row 6 is Hadir percentage
        cell_value = ws['C6'].value
        self.assertIsInstance(cell_value, str)
        self.assertTrue(cell_value.startswith('='))
    
    def test_export_excel_frozen_panes(self):
        """Test that Excel sheets have frozen panes."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date
        )
        
        # Load workbook
        wb = load_workbook(BytesIO(excel_bytes))
        
        # Check Summary sheet
        ws_summary = wb['Ringkasan']
        self.assertIsNotNone(ws_summary.freeze_panes)
        
        # Check Detail sheet
        ws_detail = wb['Detail Kehadiran']
        self.assertIsNotNone(ws_detail.freeze_panes)
    
    def test_export_excel_invalid_date_range(self):
        """Test that invalid date range raises error."""
        start_date = timezone.now().date()
        end_date = timezone.now().date() - timedelta(days=7)
        
        with self.assertRaises(TeacherReportServiceError):
            TeacherReportService.export_attendance_excel(
                start_date=start_date,
                end_date=end_date
            )
    
    def test_export_excel_contains_attendance_data(self):
        """Test that Excel export contains actual attendance data."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date
        )
        
        # Load workbook
        wb = load_workbook(BytesIO(excel_bytes))
        ws = wb['Detail Kehadiran']
        
        # Check that there's data beyond the header row
        # Row 2 should have attendance data
        self.assertIsNotNone(ws['A2'].value)  # No column
        self.assertIsNotNone(ws['B2'].value)  # Date column
        self.assertIsNotNone(ws['C2'].value)  # NIP column
    
    def test_export_excel_summary_row_with_totals(self):
        """Test that summary sheet has a total row with formulas."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        excel_bytes = TeacherReportService.export_attendance_excel(
            start_date=start_date,
            end_date=end_date
        )
        
        # Load workbook
        wb = load_workbook(BytesIO(excel_bytes))
        ws = wb['Ringkasan']
        
        # Find the summary row (should be after all teacher data)
        # Start from row 18 (first data row) and find where data ends
        row_idx = 18
        while ws.cell(row=row_idx, column=1).value is not None:
            row_idx += 1
        
        # Summary row should be at row_idx
        summary_row = row_idx
        
        # Check that TOTAL label exists
        total_cell = ws.cell(row=summary_row, column=2).value
        self.assertEqual(total_cell, 'TOTAL')
        
        # Check that formulas exist for totals
        # Column D (Total JP) should have a SUM formula
        formula_cell = ws.cell(row=summary_row, column=4).value
        if formula_cell:
            self.assertIsInstance(formula_cell, str)
            self.assertTrue(formula_cell.startswith('=SUM'))


class TeacherReportServicePDFTestCase(TestCase):
    """Test cases for PDF generation functionality."""
    
    def setUp(self):
        """Set up test data."""
        # Create admin user
        self.admin_user = User.objects.create_superuser(
            username='admin',
            email='admin@test.com',
            password='testpass123'
        )
        
        # Create academic level
        self.academic_level, _ = AcademicLevel.objects.get_or_create(
            code='SMP',
            defaults={
                'name': 'Sekolah Menengah Pertama',
                'level_type': 'SMP',
                'min_grade': 7,
                'max_grade': 9
            }
        )
        
        # Create classroom
        self.classroom = Classroom.objects.create(
            name='8A',
            academic_level=self.academic_level,
            grade=8,
            section='A',
            is_active=True
        )
        
        # Create subject
        self.subject = Subject.objects.create(
            code='MAT',
            name='Matematika',
            category='UMUM',
            is_active=True
        )
        
        # Create teacher
        self.teacher = Teacher.objects.create(
            nip='1234567890',
            full_name='Ahmad Yusuf',
            email='ahmad@test.com',
            employment_date=timezone.now().date() - timedelta(days=365),
            employment_status='ACTIVE',
            is_active=True
        )
        
        # Create schedule
        self.schedule = TeacherSchedule.objects.create(
            teacher=self.teacher,
            subject=self.subject,
            classroom=self.classroom,
            day_of_week=0,
            jp_start=1,
            jp_end=2,
            effective_date=timezone.now().date() - timedelta(days=30),
            is_active=True
        )
        
        # Create attendance records
        today = timezone.now().date()
        for i in range(5):
            date = today - timedelta(days=i)
            TeacherAttendance.objects.create(
                teacher=self.teacher,
                schedule=self.schedule,
                date=date,
                jp_number=1,
                status='HADIR',
                recorded_by=self.admin_user
            )
    
    def test_generate_pdf_basic(self):
        """Test basic PDF generation functionality."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        pdf_bytes = TeacherReportService.generate_teacher_report_pdf(
            teacher_id=self.teacher.id,
            start_date=start_date,
            end_date=end_date
        )
        
        # Verify it returns bytes
        self.assertIsInstance(pdf_bytes, bytes)
        self.assertGreater(len(pdf_bytes), 0)
        
        # Verify it's a PDF (starts with PDF magic number)
        self.assertTrue(pdf_bytes.startswith(b'%PDF'))
    
    def test_generate_pdf_invalid_teacher(self):
        """Test that invalid teacher ID raises error."""
        start_date = timezone.now().date() - timedelta(days=7)
        end_date = timezone.now().date()
        
        # Use a random UUID that doesn't exist
        from uuid import uuid4
        invalid_id = uuid4()
        
        with self.assertRaises(TeacherReportServiceError):
            TeacherReportService.generate_teacher_report_pdf(
                teacher_id=invalid_id,
                start_date=start_date,
                end_date=end_date
            )
    
    def test_generate_pdf_invalid_date_range(self):
        """Test that invalid date range raises error."""
        start_date = timezone.now().date()
        end_date = timezone.now().date() - timedelta(days=7)
        
        with self.assertRaises(TeacherReportServiceError):
            TeacherReportService.generate_teacher_report_pdf(
                teacher_id=self.teacher.id,
                start_date=start_date,
                end_date=end_date
            )
